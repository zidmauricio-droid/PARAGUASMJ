"""
routes/proyectos.py — Gestión de Proyectos SIGCA
Incluye: lista paginada, nuevo, ver, tareas, evidencias, reporte Excel.
Codificación: GA (soporte PSMV, PUEAA, SSPD)
"""
import os
import sqlite3
import logging
import time
from functools import wraps
from io import BytesIO
from datetime import date, datetime
from flask import (Blueprint, render_template, request, jsonify,
                   redirect, url_for, flash, session, send_file)
from werkzeug.utils import secure_filename
from core.database_manager import get_db, obtener_consecutivo
from core.seguridad import login_requerido, rol_requerido
from utils.audit import log_action
from utils.seguridad import verificar_token_csrf

proy_bp = Blueprint("proyectos", __name__, url_prefix="/proyectos")
logger  = logging.getLogger("sigca.proyectos")
UPLOAD_EVIDENCIAS = os.path.join("uploads", "evidencias_proyectos")
os.makedirs(UPLOAD_EVIDENCIAS, exist_ok=True)


def with_retry(max_retries: int = 3, base_delay: float = 0.25):
    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            last_exc = None
            for attempt in range(max_retries):
                try:
                    return func(*args, **kwargs)
                except sqlite3.OperationalError as e:
                    last_exc = e
                    if "database is locked" in str(e).lower() and attempt < max_retries - 1:
                        time.sleep(base_delay * (2 ** attempt)); continue
                    break
                except Exception:
                    raise
            raise last_exc
        return wrapper
    return decorator


@proy_bp.route("/")
@login_requerido
def listar():
    page     = max(request.args.get("page", 1, type=int), 1)
    per_page = min(request.args.get("per_page", 20, type=int), 100)
    offset   = (page - 1) * per_page
    conn     = get_db()
    try:
        total = conn.execute("SELECT COUNT(*) as c FROM proyectos").fetchone()["c"]
        proyectos = conn.execute("""
            SELECT p.*,
                   COUNT(t.id) as total_tareas,
                   SUM(CASE WHEN t.estado = 'completada' THEN 1 ELSE 0 END) as tareas_ok,
                   ROUND(COALESCE(AVG(t.porcentaje_avance), 0), 1) as avance_general
            FROM proyectos p
            LEFT JOIN tareas_proyecto t ON p.pk_proyecto_id = t.proyecto_id
            GROUP BY p.pk_proyecto_id
            ORDER BY p.pk_proyecto_id DESC
            LIMIT ? OFFSET ?
        """, (per_page, offset)).fetchall()
        total_pages = (total + per_page - 1) // per_page
        return render_template("proyectos/lista.html",
                               proyectos=proyectos, page=page,
                               per_page=per_page, total_pages=total_pages, total=total)
    except Exception as e:
        logger.error(f"Error listando proyectos: {e}")
        flash("Error al cargar proyectos", "danger")
        return render_template("proyectos/lista.html", proyectos=[])
    finally:
        conn.close()


@proy_bp.route("/nuevo", methods=["GET", "POST"])
@login_requerido
@rol_requerido("admin", "coordinador", "director")
def nuevo():
    if request.method == "POST":
        if not verificar_token_csrf():
            flash("Solicitud inválida. Recargue la página e intente de nuevo.", "danger")
            return redirect(url_for("proyectos.nuevo"))
        nombre = request.form.get("nombre", "").strip()
        if not nombre:
            flash("El nombre del proyecto es obligatorio.", "danger")
            return redirect(url_for("proyectos.nuevo"))
        try:
            presupuesto = float(request.form.get("presupuesto") or 0)
            if presupuesto < 0:
                flash("El presupuesto no puede ser negativo.", "danger")
                return redirect(url_for("proyectos.nuevo"))
        except (ValueError, TypeError):
            flash("El presupuesto debe ser un valor numérico.", "danger")
            return redirect(url_for("proyectos.nuevo"))
        conn = get_db()
        try:
            anio   = date.today().year
            consec = obtener_consecutivo("GA", "PRY", anio)
            codigo = f"GA-PRY-{anio}-{consec:03d}"
            responsable_id = request.form.get("responsable_id") or None
            tipo_periodo    = request.form.get("tipo_periodo", "ANIOS").upper()
            if tipo_periodo not in ("MESES", "ANIOS", "TRIMESTRES", "SEMESTRES"):
                tipo_periodo = "ANIOS"
            try:
                cantidad_periodo = max(1, int(request.form.get("cantidad_periodo") or 1))
            except (ValueError, TypeError):
                cantidad_periodo = 1
            _MESES_FACTOR = {"MESES": 1, "TRIMESTRES": 3, "SEMESTRES": 6, "ANIOS": 12}
            duracion_meses = cantidad_periodo * _MESES_FACTOR.get(tipo_periodo, 12)
            conn.execute("""
                INSERT INTO proyectos
                (codigo, nombre, descripcion, tipo_proyecto,
                 fecha_inicio, fecha_limite, presupuesto,
                 responsable_id, estado, fecha_creacion, creado_por,
                 tipo_periodo, cantidad_periodo, duracion_meses)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """, (codigo, nombre,
                  request.form.get("descripcion", "").strip(),
                  request.form.get("tipo_proyecto", "otro"),
                  request.form.get("fecha_inicio") or None,
                  request.form.get("fecha_limite") or None,
                  presupuesto,
                  responsable_id,
                  "planificacion",
                  datetime.now().isoformat(),
                  session.get("usuario_id"),
                  tipo_periodo, cantidad_periodo, duracion_meses))
            new_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
            conn.commit()
            log_action(
                accion="CREATE_PROYECTO",
                modulo="proyectos",
                descripcion=f"{codigo} — {nombre} | presupuesto={presupuesto:,.0f} | periodo={cantidad_periodo}{tipo_periodo} ({duracion_meses}m)"
            )
            flash(f"Proyecto creado: {codigo}", "success")
            return redirect(url_for("proyectos.ver", pid=new_id))
        except Exception as e:
            conn.rollback()
            logger.error(f"Error creando proyecto: {e}")
            flash("Error al crear el proyecto. Verifique los datos e intente de nuevo.", "danger")
        finally:
            conn.close()
    conn = get_db()
    responsables = conn.execute(
        "SELECT pk_usuario_id, nombre_completo FROM usuarios WHERE activo=1 ORDER BY nombre_completo"
    ).fetchall()
    conn.close()
    return render_template("proyectos/nuevo.html", responsables=responsables)


@proy_bp.route("/<int:pid>")
@login_requerido
def ver(pid):
    conn = get_db()
    try:
        proyecto = conn.execute(
            "SELECT * FROM proyectos WHERE pk_proyecto_id=?", (pid,)
        ).fetchone()
        if not proyecto:
            flash("Proyecto no encontrado", "danger")
            return redirect(url_for("proyectos.listar"))
        tareas = conn.execute("""
            SELECT t.*, u.nombre_completo as responsable_nombre
            FROM tareas_proyecto t
            LEFT JOIN usuarios u ON t.responsable_id = u.pk_usuario_id
            WHERE t.proyecto_id = ? ORDER BY t.fecha_inicio_plan, t.nombre
        """, (pid,)).fetchall()
        avance = conn.execute(
            "SELECT COALESCE(AVG(porcentaje_avance),0) as avg FROM tareas_proyecto WHERE proyecto_id=?",
            (pid,)
        ).fetchone()["avg"]
        evidencias = conn.execute("""
            SELECT e.*, u.nombre_completo as subido_por_nombre
            FROM evidencias_proyecto e
            LEFT JOIN usuarios u ON e.subido_por = u.pk_usuario_id
            WHERE e.proyecto_id = ? ORDER BY e.fecha_subida DESC
            LIMIT 50
        """, (pid,)).fetchall()
        return render_template("proyectos/ver.html",
                               proyecto=proyecto, tareas=tareas,
                               evidencias=evidencias, avance_general=round(avance, 1))
    finally:
        conn.close()


@proy_bp.route("/<int:pid>/tarea", methods=["POST"])
@login_requerido
@rol_requerido("admin", "coordinador", "director")
@with_retry()
def agregar_tarea(pid):
    csrf_ok = (
        verificar_token_csrf()
        or request.headers.get("X-CSRFToken") == session.get("csrf_token")
    )
    if not csrf_ok:
        return jsonify({"ok": False, "error": "Solicitud inválida (CSRF)"}), 403
    data = request.get_json(silent=True) or {}
    nombre_tarea = (data.get("nombre") or "").strip()
    if not nombre_tarea:
        return jsonify({"ok": False, "error": "El nombre de la tarea es obligatorio"}), 400
    try:
        costo = float(data.get("costo_estimado") or 0)
        if costo < 0:
            return jsonify({"ok": False, "error": "El costo no puede ser negativo"}), 400
    except (ValueError, TypeError):
        return jsonify({"ok": False, "error": "El costo debe ser un valor numérico"}), 400
    conn = get_db()
    try:
        conn.execute("""
            INSERT INTO tareas_proyecto
            (proyecto_id, nombre, descripcion, responsable_id,
             fecha_inicio_plan, fecha_fin_plan, costo_estimado, estado, porcentaje_avance)
            VALUES (?,?,?,?,?,?,?,'pendiente',0)
        """, (pid, nombre_tarea, (data.get("descripcion") or "").strip(),
              data.get("responsable_id"),
              data.get("fecha_inicio"), data.get("fecha_fin"),
              costo))
        conn.commit()
        log_action(
            accion="ADD_TAREA",
            modulo="proyectos",
            descripcion=f"Proyecto {pid} — tarea: {nombre_tarea} | costo={costo:,.0f}"
        )
        return jsonify({"ok": True}), 201
    except Exception as e:
        conn.rollback()
        return jsonify({"ok": False, "error": "Error interno al guardar la tarea"}), 500
    finally:
        conn.close()


@proy_bp.route("/<int:pid>/evidencia", methods=["POST"])
@login_requerido
@rol_requerido("admin", "coordinador", "director")
@with_retry()
def subir_evidencia(pid):
    if not verificar_token_csrf():
        return jsonify({"ok": False, "error": "Solicitud inválida (CSRF)"}), 403
    if "archivo" not in request.files:
        return jsonify({"ok": False, "error": "No se recibió ningún archivo"}), 400
    archivo = request.files["archivo"]
    if not archivo.filename:
        return jsonify({"ok": False, "error": "El archivo está vacío"}), 400
    ext = archivo.filename.rsplit(".", 1)[-1].lower() if "." in archivo.filename else ""
    if ext not in {"pdf", "jpg", "jpeg", "png", "doc", "docx", "xls", "xlsx"}:
        return jsonify({"ok": False, "error": f"Formato no permitido: {ext}. Use PDF, imágenes o documentos Office"}), 400
    fn   = secure_filename(f"PRY-{pid}-{datetime.now().strftime('%Y%m%d%H%M%S')}.{ext}")
    ruta = os.path.join(UPLOAD_EVIDENCIAS, fn)
    conn = get_db()
    try:
        conn.execute("""
            INSERT INTO evidencias_proyecto
            (proyecto_id, nombre_archivo, ruta, descripcion, subido_por, fecha_subida)
            VALUES (?,?,?,?,?,datetime('now'))
        """, (pid, fn, ruta, (request.form.get("descripcion") or ""), session.get("usuario_id")))
        archivo.save(ruta)
        conn.commit()
        log_action(
            accion="UPLOAD_EVIDENCIA",
            modulo="proyectos",
            descripcion=f"Proyecto {pid} — archivo: {fn}"
        )
        return jsonify({"ok": True}), 201
    except Exception as e:
        conn.rollback()
        if os.path.exists(ruta):
            try:
                os.remove(ruta)
            except OSError:
                pass
        logger.error(f"Error subiendo evidencia proyecto {pid}: {e}")
        return jsonify({"ok": False, "error": "Error interno al guardar la evidencia"}), 500
    finally:
        conn.close()


@proy_bp.route("/<int:pid>/reporte")
@login_requerido
def reporte_proyecto(pid):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        flash("openpyxl no instalado. Ejecute: pip install openpyxl", "danger")
        return redirect(url_for("proyectos.ver", pid=pid))
    conn = get_db()
    try:
        proyecto = conn.execute(
            "SELECT * FROM proyectos WHERE pk_proyecto_id=?", (pid,)
        ).fetchone()
        if not proyecto:
            flash("Proyecto no encontrado","danger")
            return redirect(url_for("proyectos.listar"))
        tareas = conn.execute("""
            SELECT t.*, u.nombre_completo as responsable
            FROM tareas_proyecto t
            LEFT JOIN usuarios u ON t.responsable_id = u.pk_usuario_id
            WHERE t.proyecto_id=? ORDER BY t.fecha_inicio_plan
        """, (pid,)).fetchall()
        wb  = Workbook()
        ws  = wb.active
        ws.title = "Proyecto"
        ws["A1"] = f"{proyecto['codigo']} — {proyecto['nombre']}"
        ws["A1"].font = Font(size=13, bold=True, color="1E3A8A")
        ws.merge_cells("A1:F1")
        azul = PatternFill("solid", fgColor="1E3A8A")
        hdrs = ["Tarea","Responsable","Inicio","Fin","Avance%","Estado"]
        for col, h in enumerate(hdrs,1):
            c = ws.cell(2, col, h)
            c.font = Font(bold=True, color="FFFFFF"); c.fill = azul
            c.alignment = Alignment(horizontal="center")
        for i,t in enumerate(tareas,3):
            ws.cell(i,1,t["nombre"])
            ws.cell(i,2,t.get("responsable",""))
            ws.cell(i,3,t.get("fecha_inicio_plan",""))
            ws.cell(i,4,t.get("fecha_fin_plan",""))
            ws.cell(i,5,t.get("porcentaje_avance",0))
            ws.cell(i,6,t.get("estado",""))
        output = BytesIO(); wb.save(output); output.seek(0)
        fn = f"Proyecto_{proyecto['codigo']}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return send_file(output, as_attachment=True, download_name=fn,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        logger.error(f"Error reporte proyecto {pid}: {e}")
        flash("Error generando reporte","danger")
        return redirect(url_for("proyectos.ver", pid=pid))
    finally:
        conn.close()


print("✅ Módulo proyectos cargado")
